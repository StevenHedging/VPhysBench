from __future__ import annotations

from pathlib import Path
import tempfile
import unittest
import zipfile

from openpyxl import Workbook

from physbench.datasets.vertical_spring_import import (
    SourceMember,
    TrialAnnotation,
    inventory_archive,
    load_trial_annotations,
    map_trials_to_sources,
)


def _write_workbook(path: Path, rows: list[tuple[object, ...]]) -> Path:
    workbook = Workbook()
    settings = workbook.active
    settings.title = "实验设置"
    spring = workbook.create_sheet("弹簧设置")
    spring.append(["弹簧ID", "S01"])
    trials = workbook.create_sheet("Trial记录")
    trials.append(["Trial ID", "弹簧ID", "视频文件", "初始位移 (mm)"])
    for row in rows:
        trials.append(row)
    workbook.save(path)
    return path


class IntakeTests(unittest.TestCase):
    def test_load_trials_discovers_real_header_row_and_adds_mov_suffix(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "spring.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Trial记录"
            sheet.append(["Trial 记录｜每次释放一行"])
            sheet.append([])
            sheet.append(["整批统一静止释放。"])
            sheet.append([])
            sheet.append([
                "Trial编号",
                "弹簧编号",
                "重复编号",
                "视频文件名",
                "释放长度-平衡长度_mm",
            ])
            sheet.append(["T001", "S01", None, "IMG_1538", 10])
            workbook.save(path)

            trials = load_trial_annotations(path)

        self.assertEqual("IMG_1538.MOV", trials[0].video_name)
        self.assertEqual(6, trials[0].workbook_row)

    def test_load_trials_preserves_signed_displacement(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workbook = _write_workbook(
                Path(directory) / "spring.xlsx",
                [
                    ("T001", "S01", "IMG_1518.MOV", 30.0),
                    ("T002", "S01", "IMG_1519.MOV", -40.0),
                    ("T003", "S01", None, None),
                ],
            )

            trials = load_trial_annotations(workbook)

        self.assertEqual(
            [
                ("T001", 30.0, "below", 0.03, 2),
                ("T002", -40.0, "above", 0.04, 3),
            ],
            [
                (
                    item.trial_id,
                    item.displacement_mm,
                    item.direction,
                    item.displacement_m,
                    item.workbook_row,
                )
                for item in trials
            ],
        )

    def test_mapping_excludes_missing_required_video(self) -> None:
        trial = TrialAnnotation("T001", "S01", "IMG_1652.MOV", 30.0, 2)

        result = map_trials_to_sources([trial], [])

        self.assertEqual((), result.accepted)
        self.assertEqual(1, len(result.exclusions))
        self.assertEqual(
            "workbook_video_missing_from_archive",
            result.exclusions[0].reason,
        )

    def test_inventory_uses_zip_crc_and_basename(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            archive = Path(directory) / "source.zip"
            with zipfile.ZipFile(archive, "w") as handle:
                handle.writestr("batch/IMG_1518.MOV", b"video-a")
                handle.writestr("batch/notes.xlsx", b"workbook")

            members = inventory_archive(archive)

        self.assertEqual(1, len(members))
        self.assertEqual("IMG_1518.MOV", members[0].basename)
        self.assertEqual(len(b"video-a"), members[0].size)
        self.assertIsInstance(members[0].crc32, int)

    def test_mapping_rejects_nonidentical_duplicate_basenames(self) -> None:
        trial = TrialAnnotation("T001", "S01", "IMG_1518.MOV", 30.0, 2)
        sources = [
            SourceMember("a/IMG_1518.MOV", "IMG_1518.MOV", 10, 1),
            SourceMember("b/IMG_1518.MOV", "IMG_1518.MOV", 11, 2),
        ]

        result = map_trials_to_sources([trial], sources)

        self.assertEqual((), result.accepted)
        self.assertEqual("ambiguous_source_members", result.exclusions[0].reason)

    def test_load_trials_rejects_zero_displacement(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workbook = _write_workbook(
                Path(directory) / "spring.xlsx",
                [("T001", "S01", "IMG_1518.MOV", 0.0)],
            )

            with self.assertRaisesRegex(ValueError, "invalid displacement"):
                load_trial_annotations(workbook)

    def test_load_trials_rejects_unsupported_spring(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workbook = _write_workbook(
                Path(directory) / "spring.xlsx",
                [("T001", "S02", "IMG_1518.MOV", 30.0)],
            )

            with self.assertRaisesRegex(ValueError, "unsupported spring ID"):
                load_trial_annotations(workbook)


if __name__ == "__main__":
    unittest.main()
