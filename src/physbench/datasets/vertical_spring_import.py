from __future__ import annotations

from dataclasses import dataclass
import math
from pathlib import Path, PurePosixPath
from typing import Literal, Sequence
import zipfile

from openpyxl import load_workbook


Direction = Literal["above", "below"]


@dataclass(frozen=True)
class TrialAnnotation:
    trial_id: str
    spring_id: str
    video_name: str
    displacement_mm: float
    workbook_row: int

    @property
    def direction(self) -> Direction:
        return "below" if self.displacement_mm > 0 else "above"

    @property
    def displacement_m(self) -> float:
        return abs(self.displacement_mm) / 1000.0


@dataclass(frozen=True)
class SourceMember:
    member: str
    basename: str
    size: int
    crc32: int


@dataclass(frozen=True)
class MappedTrial:
    trial: TrialAnnotation
    source: SourceMember
    duplicate_members: tuple[str, ...] = ()


@dataclass(frozen=True)
class IntakeExclusion:
    trial_id: str
    video_name: str
    reason: str
    details: tuple[str, ...] = ()


@dataclass(frozen=True)
class IntakeResult:
    accepted: tuple[MappedTrial, ...]
    exclusions: tuple[IntakeExclusion, ...]


def _normalized_header(value: object) -> str:
    return "".join(str(value or "").strip().lower().split())


def _column_index(headers: Sequence[object], *names: str) -> int:
    normalized = [_normalized_header(value) for value in headers]
    for name in names:
        target = _normalized_header(name)
        if target in normalized:
            return normalized.index(target)
    raise ValueError(f"missing required workbook column: {names[0]}")


def load_trial_annotations(workbook: Path) -> list[TrialAnnotation]:
    document = load_workbook(workbook, data_only=True, read_only=True)
    if "Trial记录" not in document.sheetnames:
        raise ValueError("workbook has no Trial记录 sheet")
    sheet = document["Trial记录"]
    rows = list(sheet.iter_rows(values_only=True))
    header_offset = next(
        (
            index
            for index, row in enumerate(rows)
            if _normalized_header("Trial编号")
            in {_normalized_header(value) for value in row}
            or _normalized_header("Trial ID")
            in {_normalized_header(value) for value in row}
        ),
        None,
    )
    if header_offset is None:
        raise ValueError("Trial记录 sheet has no recognized header row")
    headers = rows[header_offset]
    trial_col = _column_index(headers, "Trial ID", "TrialID", "Trial编号")
    spring_col = _column_index(headers, "弹簧ID", "Spring ID", "弹簧编号")
    video_col = _column_index(
        headers,
        "视频文件",
        "视频名",
        "Video",
        "视频文件名",
    )
    displacement_col = _column_index(
        headers,
        "初始位移 (mm)",
        "初始位移(mm)",
        "位移(mm)",
        "释放长度-平衡长度_mm",
    )
    result: list[TrialAnnotation] = []
    for workbook_row, row in enumerate(
        rows[header_offset + 1 :],
        start=header_offset + 2,
    ):
        trial_id = str(row[trial_col] or "").strip()
        spring_id = str(row[spring_col] or "").strip()
        video_name = str(row[video_col] or "").strip()
        displacement = row[displacement_col]
        if not trial_id:
            continue
        if not video_name and displacement in (None, ""):
            continue
        if not spring_id or not video_name or displacement in (None, ""):
            raise ValueError(f"incomplete required Trial row {workbook_row}")
        if spring_id != "S01":
            raise ValueError(
                f"unsupported spring ID {spring_id!r} at row {workbook_row}"
            )
        displacement_mm = float(displacement)
        if not math.isfinite(displacement_mm) or displacement_mm == 0:
            raise ValueError(f"invalid displacement at row {workbook_row}")
        normalized_video_name = PurePosixPath(video_name).name
        if not PurePosixPath(normalized_video_name).suffix:
            normalized_video_name = f"{normalized_video_name}.MOV"
        result.append(
            TrialAnnotation(
                trial_id=trial_id,
                spring_id=spring_id,
                video_name=normalized_video_name,
                displacement_mm=displacement_mm,
                workbook_row=workbook_row,
            )
        )
    return result


def inventory_archive(archive: Path) -> list[SourceMember]:
    with zipfile.ZipFile(archive) as handle:
        members = [
            SourceMember(
                member=info.filename,
                basename=PurePosixPath(info.filename).name,
                size=info.file_size,
                crc32=info.CRC,
            )
            for info in handle.infolist()
            if not info.is_dir() and info.filename.lower().endswith(".mov")
        ]
    return sorted(members, key=lambda item: (item.basename.lower(), item.member))


def map_trials_to_sources(
    trials: Sequence[TrialAnnotation],
    sources: Sequence[SourceMember],
) -> IntakeResult:
    sources_by_name: dict[str, list[SourceMember]] = {}
    for source in sources:
        sources_by_name.setdefault(source.basename.casefold(), []).append(source)
    accepted: list[MappedTrial] = []
    exclusions: list[IntakeExclusion] = []
    for trial in trials:
        matches = sources_by_name.get(trial.video_name.casefold(), [])
        if not matches:
            exclusions.append(
                IntakeExclusion(
                    trial.trial_id,
                    trial.video_name,
                    "workbook_video_missing_from_archive",
                )
            )
            continue
        signatures = {(item.size, item.crc32) for item in matches}
        if len(signatures) != 1:
            exclusions.append(
                IntakeExclusion(
                    trial.trial_id,
                    trial.video_name,
                    "ambiguous_source_members",
                    tuple(item.member for item in matches),
                )
            )
            continue
        selected = sorted(matches, key=lambda item: item.member)[0]
        accepted.append(
            MappedTrial(
                trial=trial,
                source=selected,
                duplicate_members=tuple(
                    item.member for item in matches if item.member != selected.member
                ),
            )
        )
    return IntakeResult(tuple(accepted), tuple(exclusions))
